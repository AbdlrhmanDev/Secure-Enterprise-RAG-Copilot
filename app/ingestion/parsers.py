"""Parsers for PDF, DOCX, Markdown, HTML and TXT. Each returns ordered text blocks with page/section metadata."""

from __future__ import annotations

import io
import math
import re
from dataclasses import dataclass, field
from pathlib import Path

from app.ingestion.cleaner import strip_repeated_lines
from app.text import normalize_text, words

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".markdown", ".html", ".htm", ".txt"}
WORDS_PER_PAGE = 500  # used to estimate page counts for formats without real pages


class ParseError(Exception):
    """The file cannot be parsed. Not retryable: retrying the same bytes fails the same way."""


@dataclass
class Block:
    text: str
    page: int | None = None
    section: str | None = None


@dataclass
class ParsedDocument:
    blocks: list[Block]
    page_count: int
    title: str | None = None
    meta: dict = field(default_factory=dict)


def parse_document(filename: str, data: bytes) -> ParsedDocument:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ParseError(f"unsupported file type: {ext or '(none)'}")
    try:
        if ext == ".pdf":
            parsed = _parse_pdf(data)
        elif ext == ".docx":
            parsed = _parse_docx(data)
        elif ext in (".html", ".htm"):
            parsed = _parse_html(_decode(data))
        elif ext in (".md", ".markdown"):
            parsed = _parse_markdown(_decode(data))
        else:
            parsed = _parse_txt(_decode(data))
    except ParseError:
        raise
    except Exception as exc:
        # Parser libraries raise a zoo of exception types on corrupt input.
        raise ParseError(f"could not parse {ext} file: {type(exc).__name__}") from exc

    parsed.blocks = [b for b in parsed.blocks if b.text]
    if not parsed.blocks:
        raise ParseError("document contains no extractable text")
    if not parsed.page_count:
        total_words = sum(len(words(b.text)) for b in parsed.blocks)
        parsed.page_count = max(1, math.ceil(total_words / WORDS_PER_PAGE))
        parsed.meta["page_count_estimated"] = True
    parsed.title = parsed.title or Path(filename).stem.replace("-", " ").replace("_", " ").title()
    return parsed


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeError:
            continue
    return data.decode("latin-1")


# --- Markdown -----------------------------------------------------------------------------

_MD_HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_MD_BULLET = re.compile(r"^\s*(?:[-*+]|\d+[.)])\s+")


def _clean_inline_markdown(text: str) -> str:
    text = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"(\*\*|__|\*|`)", "", text)
    return normalize_text(text)


def _parse_markdown(text: str) -> ParsedDocument:
    blocks: list[Block] = []
    title: str | None = None
    section: str | None = None
    buffer: list[str] = []
    in_code = False

    def flush() -> None:
        if buffer:
            blocks.append(Block(_clean_inline_markdown(" ".join(buffer)), None, section))
            buffer.clear()

    for line in text.splitlines():
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        heading = None if in_code else _MD_HEADING.match(line)
        if heading:
            flush()
            section = _clean_inline_markdown(heading.group(2))
            if len(heading.group(1)) == 1 and title is None:
                title = section
        elif not line.strip():
            flush()
        else:
            buffer.append(line.strip() if in_code else _MD_BULLET.sub("", line).strip())
    flush()
    return ParsedDocument(blocks, 0, title)


# --- Plain text ---------------------------------------------------------------------------

_UNDERLINE = re.compile(r"^\s*(={3,}|-{3,})\s*$")


def _parse_txt(text: str) -> ParsedDocument:
    blocks: list[Block] = []
    title: str | None = None
    section: str | None = None
    buffer: list[str] = []
    lines = text.splitlines()

    def flush() -> None:
        if buffer:
            blocks.append(Block(normalize_text(" ".join(buffer)), None, section))
            buffer.clear()

    i = 0
    while i < len(lines):
        line = lines[i].strip()
        underlined = i + 1 < len(lines) and _UNDERLINE.match(lines[i + 1]) and line
        shouting = line.isupper() and len(line.split()) <= 10 and not buffer and any(c.isalpha() for c in line)
        if underlined or shouting:
            flush()
            section = normalize_text(line)
            title = title or section
            i += 2 if underlined else 1
            continue
        if not line or _UNDERLINE.match(line):
            flush()
        else:
            buffer.append(line)
        i += 1
    flush()
    return ParsedDocument(blocks, 0, title)


# --- HTML ---------------------------------------------------------------------------------

_HTML_HEADINGS = ("h1", "h2", "h3", "h4", "h5", "h6")
_HTML_CONTENT = ("p", "li", "pre", "tr", "blockquote", "dd", "dt")


def _parse_html(text: str) -> ParsedDocument:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(text, "html.parser")
    for tag in soup(["script", "style", "nav", "header", "footer", "aside", "noscript", "form"]):
        tag.decompose()
    title = normalize_text(soup.title.get_text()) if soup.title else None

    blocks: list[Block] = []
    section: str | None = None
    for el in soup.find_all([*_HTML_HEADINGS, *_HTML_CONTENT]):
        if el.find_parent(_HTML_CONTENT):
            continue  # nested content is already covered by its ancestor's text
        content = normalize_text(el.get_text(" ", strip=True))
        if not content:
            continue
        if el.name in _HTML_HEADINGS:
            section = content
            title = title or content
        else:
            blocks.append(Block(content, None, section))
    if not blocks:
        body = normalize_text(soup.get_text(" ", strip=True))
        blocks = [Block(body, None, None)] if body else []
    return ParsedDocument(blocks, 0, title)


# --- DOCX ---------------------------------------------------------------------------------


def _parse_docx(data: bytes) -> ParsedDocument:
    import docx
    from docx.table import Table

    document = docx.Document(io.BytesIO(data))
    blocks: list[Block] = []
    title = normalize_text(document.core_properties.title or "") or None
    section: str | None = None

    for item in document.iter_inner_content():
        if isinstance(item, Table):
            for row in item.rows:
                cells = [normalize_text(c.text) for c in row.cells]
                content = " | ".join(c for c in cells if c)
                if content:
                    blocks.append(Block(content, None, section))
            continue
        content = normalize_text(item.text)
        if not content:
            continue
        style = (item.style.name or "") if item.style is not None else ""
        if style.startswith(("Heading", "Title")):
            section = content
            title = title or content
        else:
            blocks.append(Block(content, None, section))
    return ParsedDocument(blocks, 0, title)


# --- PDF ----------------------------------------------------------------------------------

# "1. Scope", "2.1. Details" or "2.1 Details". A bare "25 USD ..." line is body text, not a heading.
_PDF_NUMBERED = re.compile(r"^(?:\d+(?:\.\d+)*\.|\d+(?:\.\d+)+)\s+[A-Z][^.]{0,78}$")
_PAGE_NUMBER = re.compile(r"^(?:page\s+)?\d+(?:\s*(?:of|/)\s*\d+)?$", re.IGNORECASE)


def _is_pdf_heading(line: str) -> bool:
    if len(line) > 80 or line.endswith((".", ",", ";", ":")):
        return False
    if _PDF_NUMBERED.match(line):
        return True
    return line.isupper() and 1 <= len(line.split()) <= 10 and sum(c.isalpha() for c in line) >= 4


def _parse_pdf(data: bytes) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise ParseError("PDF is encrypted")

    pages = []
    for page in reader.pages:
        raw = page.extract_text() or ""
        lines = [line.strip() for line in raw.splitlines() if line.strip()]
        pages.append([line for line in lines if not _PAGE_NUMBER.match(line)])
    pages = strip_repeated_lines(pages)

    blocks: list[Block] = []
    section: str | None = None
    for page_number, lines in enumerate(pages, start=1):
        buffer: list[str] = []
        for line in lines:
            if _is_pdf_heading(line):
                if buffer:
                    blocks.append(Block(normalize_text(" ".join(buffer)), page_number, section))
                    buffer = []
                section = normalize_text(re.sub(r"^[\d.]+\s+", "", line))
            else:
                buffer.append(line)
        if buffer:
            blocks.append(Block(normalize_text(" ".join(buffer)), page_number, section))

    info_title = None
    try:
        info_title = normalize_text(reader.metadata.title or "") if reader.metadata else None
    except Exception:
        info_title = None
    # The title line printed on page one is metadata, not a passage worth indexing on its own.
    if info_title and blocks and blocks[0].text.endswith(info_title) and len(blocks[0].text) < len(info_title) + 80:
        blocks.pop(0)
    return ParsedDocument(blocks, len(reader.pages), info_title or None)
