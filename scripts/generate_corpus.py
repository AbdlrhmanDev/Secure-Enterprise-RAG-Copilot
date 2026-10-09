"""Generate the demo corpus and the labeled evaluation dataset.

    python -m scripts.generate_corpus                 # core documents + ~100 filler documents (500+ pages)
    python -m scripts.generate_corpus --filler-docs 0 # core documents only

Outputs:
    data/corpus/<collection>/<file>   documents in PDF, DOCX, Markdown, HTML and TXT
    data/corpus/manifest.json         collection + ACL for every file (read by scripts.ingest_corpus)
    evals/dataset.jsonl               labeled evaluation cases

The core documents carry every fact the evaluation asks about. The filler documents are
realistic-looking distractors (meeting reviews, release notes, vendor assessments, status
reports) that add volume and retrieval noise; no evaluation question is answered by them.
"""

from __future__ import annotations

import argparse
import html
import json
import random
from pathlib import Path

from scripts.corpus_data import ASKER_BY_ROLE, DOCS, OUTSIDERS_BY_ROLE, PARAPHRASES, UNANSWERABLE

CORPUS_DIR = Path("data/corpus")
DATASET_PATH = Path("evals/dataset.jsonl")
COMPANY = "Halden Systems"

Section = tuple[str, str]  # (heading, body with paragraphs separated by blank lines)


# --- Renderers ------------------------------------------------------------------------------


def render_markdown(title: str, sections: list[Section]) -> bytes:
    parts = [f"# {title}", ""]
    for heading, body in sections:
        parts += [f"## {heading}", "", body, ""]
    return "\n".join(parts).encode()


def render_txt(title: str, sections: list[Section]) -> bytes:
    parts = [title, "=" * len(title), ""]
    for heading, body in sections:
        parts += [heading, "-" * len(heading), "", body, ""]
    return "\n".join(parts).encode()


def render_html(title: str, sections: list[Section]) -> bytes:
    body = []
    for heading, text in sections:
        body.append(f"<h2>{html.escape(heading)}</h2>")
        body += [f"<p>{html.escape(p)}</p>" for p in text.split("\n\n")]
    page = (
        f"<!doctype html><html><head><title>{html.escape(title)}</title><style>body{{font-family:sans-serif}}</style>"
        f"</head><body><nav><a href='/'>Intranet home</a> | <a href='/docs'>All documents</a></nav>"
        f"<h1>{html.escape(title)}</h1>{''.join(body)}"
        f"<footer>{COMPANY} internal. Do not distribute.</footer></body></html>"
    )
    return page.encode()


def render_docx(title: str, sections: list[Section], path: Path) -> None:
    import docx

    document = docx.Document()
    document.core_properties.title = title
    document.add_heading(title, level=0)
    for heading, body in sections:
        document.add_heading(heading, level=1)
        for paragraph in body.split("\n\n"):
            document.add_paragraph(paragraph)
    document.save(str(path))


def render_pdf(title: str, sections: list[Section], path: Path) -> None:
    from fpdf import FPDF

    class Pdf(FPDF):
        def header(self) -> None:
            self.set_font("Helvetica", "I", 8)
            self.cell(0, 8, f"{COMPANY} - Internal Document", align="C", new_x="LMARGIN", new_y="NEXT")

        def footer(self) -> None:
            self.set_y(-15)
            self.set_font("Helvetica", "I", 8)
            self.cell(0, 10, f"Page {self.page_no()}", align="C")

    pdf = Pdf()
    pdf.set_title(title)
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 10, title, new_x="LMARGIN", new_y="NEXT")
    for number, (heading, body) in enumerate(sections, start=1):
        pdf.ln(2)
        pdf.set_font("Helvetica", "B", 12)
        pdf.multi_cell(0, 8, f"{number}. {heading}", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 11)
        for paragraph in body.split("\n\n"):
            pdf.multi_cell(0, 6, paragraph, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)
    pdf.output(str(path))


def write_document(path: Path, title: str, sections: list[Section]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    suffix = path.suffix
    if suffix == ".pdf":
        render_pdf(title, sections, path)
    elif suffix == ".docx":
        render_docx(title, sections, path)
    else:
        renderer = {".md": render_markdown, ".txt": render_txt, ".html": render_html}[suffix]
        path.write_bytes(renderer(title, sections))


# --- Filler documents -----------------------------------------------------------------------

TEAMS = ["Atlas", "Borealis", "Cinder", "Drift", "Ember", "Fathom", "Garnet", "Harbor", "Indigo", "Juniper", "Kestrel", "Lumen"]
SYSTEMS = ["billing gateway", "search indexer", "notification relay", "partner portal", "reporting warehouse", "mobile client",
           "identity bridge", "inventory sync", "pricing engine", "document store", "audit trail", "scheduling service"]
PEOPLE = ["Amara", "Bastien", "Chiara", "Dmitri", "Elif", "Farah", "Gustav", "Hana", "Ines", "Jonas", "Keiko", "Lars"]
VENDORS = ["Northgate Data", "Pellucid Labs", "Quillon Cloud", "Riverbend Analytics", "Sablewood Software", "Tarn Networks"]
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]
ASPECTS = ["Progress", "Risks", "Dependencies", "Capacity", "Quality", "Customer Feedback", "Open Questions", "Next Steps", "Decisions"]

FILLER_KINDS = {
    # kind: (collection, roles, title template, sentence templates)
    "weekly-review": ("engineering", ["engineering"], "{team} Team Weekly Review, Week {n}", [
        "The {team} team reviewed the {system} backlog and agreed to prioritise {n} items for the coming sprint.",
        "{person} reported that the {system} migration is {pct} percent complete and remains on track for {month}.",
        "Latency on the {system} improved after {person} removed a redundant call to the {system2}.",
        "The team discussed whether the {system} should own retries or delegate them to the {system2}.",
        "{person} will pair with {person2} to investigate the flaky integration tests in the {system}.",
        "A spike of {n} failed jobs in the {system} was traced back to a configuration change made in {month}.",
        "The {team} team agreed to add a dashboard for the {system} before the next planning session.",
        "{person} raised a concern that the {system} and the {system2} use different naming conventions.",
        "Work on the {system} was paused for two days while the {team} team supported the {team2} team.",
        "The review closed with {n} action items, most of them assigned to {person} and {person2}.",
    ]),
    "release-notes": ("engineering", ["employee"], "{system} Release Notes, Build {n}", [
        "This build of the {system} adds an export option requested by the {team} team.",
        "A defect that caused duplicate entries in the {system} when the {system2} was unavailable has been fixed.",
        "The settings page of the {system} now loads {pct} percent faster on slow connections.",
        "{person} contributed a change that lets the {system} resume interrupted uploads.",
        "Support for the legacy import format in the {system} will be removed in the {month} release.",
        "The {system} now shows a clearer message when a request to the {system2} times out.",
        "Accessibility of the {system} was improved with better keyboard navigation and labels.",
        "A total of {n} minor issues reported by the {team} team were resolved in this build.",
        "The {system} documentation was updated by {person} to describe the new filter options.",
        "Known issue: sorting in the {system} is slow for lists with more than {n} thousand rows.",
    ]),
    "vendor-assessment": ("policies", ["finance"], "Vendor Assessment: {vendor}, {month} Review", [
        "{vendor} was assessed by {person} against the standard commercial and delivery criteria.",
        "The proposal from {vendor} covers support for the {system} over a term agreed with the {team} team.",
        "{person} noted that {vendor} answered {pct} percent of the questionnaire without follow up.",
        "References provided by {vendor} described delivery as dependable, with {n} open issues at handover.",
        "The {team} team compared {vendor} with {vendor2} and found the onboarding effort to be similar.",
        "{vendor} proposed a phased rollout that begins with the {system} and later extends to the {system2}.",
        "Commercial terms offered by {vendor} were reviewed by {person2} and returned with {n} comments.",
        "The assessment panel asked {vendor} to clarify how it would staff the {month} transition.",
        "{person} recommended a pilot with {vendor} limited to the {system} before any wider commitment.",
        "A follow up session with {vendor} is planned for {month} to close the remaining questions.",
    ]),
    "project-status": ("policies", ["employee"], "Project {team} Status Report, {month}", [
        "Project {team} completed {n} of its planned milestones during {month}.",
        "{person} confirmed that training material for the {system} is ready for review.",
        "The rollout of the {system} to the next office was moved to {month} at the request of the {team2} team.",
        "Feedback from the pilot group was positive, with {pct} percent rating the {system} as easy to use.",
        "{person} and {person2} ran three workshops to collect requirements for the {system2}.",
        "The steering group asked Project {team} to publish a short summary after each phase.",
        "A dependency on the {system2} remains the main schedule risk for Project {team}.",
        "Communication to affected teams will be sent by {person} once the {month} date is confirmed.",
        "Project {team} logged {n} change requests, most of them related to the {system}.",
        "The next status report will cover adoption of the {system} across the {team2} team.",
    ]),
}
FILLER_FORMATS = [".pdf", ".txt", ".md", ".html", ".docx"]


def build_filler(index: int, rng: random.Random) -> tuple[str, str, list[str], str, list[Section]]:
    kind = list(FILLER_KINDS)[index % len(FILLER_KINDS)]
    collection, roles, title_template, templates = FILLER_KINDS[kind]

    def fill(template: str) -> str:
        team, team2 = rng.sample(TEAMS, 2)
        system, system2 = rng.sample(SYSTEMS, 2)
        person, person2 = rng.sample(PEOPLE, 2)
        vendor, vendor2 = rng.sample(VENDORS, 2)
        return template.format(
            team=team, team2=team2, system=system, system2=system2, person=person, person2=person2,
            vendor=vendor, vendor2=vendor2, month=rng.choice(MONTHS), n=rng.randint(2, 48), pct=rng.randint(35, 98),
        )

    title = fill(title_template)
    title = title[0].upper() + title[1:]
    sections = []
    for aspect in ASPECTS:
        paragraphs = [" ".join(fill(rng.choice(templates)) for _ in range(5)) for _ in range(4)]
        sections.append((aspect, "\n\n".join(paragraphs)))
    filename = f"{kind}-{index:03d}{FILLER_FORMATS[index % len(FILLER_FORMATS)]}"
    return filename, collection, roles, title, sections


# --- Main -----------------------------------------------------------------------------------


def build_dataset() -> list[dict]:
    cases: list[dict] = []

    def add(query: str, expected: str, evidence: list[dict], tags: list[str], collection: str, user: str) -> None:
        cases.append({
            "id": f"q{len(cases) + 1:03d}", "query": query, "expected_answer": expected, "evidence": evidence,
            "tags": tags, "collection_id": collection, "user_id": user,
        })

    acl_cases = []
    seen: set[str] = set()
    for filename, collection, roles, _title, sections in DOCS:
        first = True
        for _heading, body, qas in sections:
            for question, answer, quote in qas:
                if quote not in body:
                    raise ValueError(f"{filename}: evidence quote is not in the section body: {quote!r}")
                evidence = [{"source": filename, "quote": quote}]
                query = PARAPHRASES.get(question, question)
                seen.add(question)
                style = "paraphrase" if query != question else "literal"
                add(query, answer, evidence, ["answerable", style, f"role:{roles[0]}"], collection, ASKER_BY_ROLE[roles[0]])
                if first:
                    for outsider in OUTSIDERS_BY_ROLE.get(roles[0], []):
                        acl_cases.append((question, answer, evidence, ["acl", f"role:{roles[0]}"], collection, outsider))
                    first = False
    unknown = set(PARAPHRASES) - seen
    if unknown:
        raise ValueError(f"paraphrases for unknown questions: {sorted(unknown)}")
    for question, collection, user in UNANSWERABLE:
        add(question, "", [], ["unanswerable"], collection, user)
    for case in acl_cases:
        add(*case)
    return cases


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--filler-docs", type=int, default=100, help="number of distractor documents (default 100)")
    parser.add_argument("--seed", type=int, default=7)
    args = parser.parse_args()

    manifest = []
    for filename, collection, roles, title, sections in DOCS:
        purpose = (
            f"This document describes the {title} of {COMPANY}. It applies to everyone on the access list of "
            "this document and is maintained by the owning team, which reviews it once a year."
        )
        rendered = [("Purpose and Scope", purpose)] + [(heading, body) for heading, body, _ in sections]
        path = CORPUS_DIR / collection / filename
        write_document(path, title, rendered)
        manifest.append({"path": path.as_posix(), "collection_id": collection, "allowed_roles": roles, "kind": "core"})

    rng = random.Random(args.seed)
    for index in range(args.filler_docs):
        filename, collection, roles, title, sections = build_filler(index, rng)
        path = CORPUS_DIR / collection / filename
        write_document(path, title, sections)
        manifest.append({"path": path.as_posix(), "collection_id": collection, "allowed_roles": roles, "kind": "filler"})

    (CORPUS_DIR / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    cases = build_dataset()
    DATASET_PATH.parent.mkdir(parents=True, exist_ok=True)
    with DATASET_PATH.open("w", encoding="utf-8", newline="\n") as handle:
        for case in cases:
            handle.write(json.dumps(case) + "\n")

    by_tag = {tag: sum(tag in c["tags"] for c in cases) for tag in ("answerable", "paraphrase", "unanswerable", "acl")}
    print(f"wrote {len(manifest)} documents to {CORPUS_DIR} ({len(DOCS)} core, {args.filler_docs} filler)")
    print(f"wrote {len(cases)} evaluation cases to {DATASET_PATH}: {by_tag}")


if __name__ == "__main__":
    main()
